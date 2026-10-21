"""Turn an uploaded document (.docx or scanned image) into plain text, then into tokens."""
from collections import namedtuple
from pathlib import Path

Token = namedtuple("Token", "text line col")   # col = index of the token inside its line
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


def read_docx(path):
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" ".join(c.text for c in row.cells))
    return "\n".join(parts)


def read_image(path):
    import pytesseract
    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    from PIL import Image
    img = Image.open(str(path)).convert("L")
    if min(img.size) < 1200:                       # tiny scans OCR badly; upscale
        f = 1200 / min(img.size)
        img = img.resize((int(img.width * f), int(img.height * f)))
    return pytesseract.image_to_string(img)


def read_document(path):
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        return read_docx(path)
    if ext in IMAGE_EXT:
        return read_image(path)
    raise ValueError(f"unsupported file type: {ext} (use .docx or an image)")


def tokenize(text):
    """Whitespace tokens that remember their line / column (layout is a useful signal)."""
    toks = []
    for li, line in enumerate(text.splitlines()):
        for ci, w in enumerate(line.split()):
            toks.append(Token(w, li, ci))
    return toks
